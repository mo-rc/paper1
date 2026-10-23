"""Build the Gate 2 hand-check workbook for all 14 cases (8 required, 6 optional) (an Excel file with live formulas).

Usage (repository root):  python scripts/make_hand_check_workbook.py [output.xlsx]
Needs data/raw/ with the section properties workbook and the two Blue Book resistance workbooks, and data/gate2/cases.csv.

Each case sheet holds: inputs typed from the Blue Book properties table; Blue Book table values (Mc, Vc, FRd, Mb);
the hand calculation as Excel formulas; the program's value for every quantity (copied from the checker);
yellow cells for the engineer's own value; differences and status computed automatically.
The workbook contains Blue Book data and is not committed to the repository."""
import csv
import os
import sys
from dataclasses import replace

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "src"))
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from ub_rl import bluebook as bb  # noqa: E402
from ub_rl.catalogue import parse_workbook  # noqa: E402
from ub_rl.checker import Checker  # noqa: E402
from ub_rl.config import CheckerConfig, PHI_REF  # noqa: E402

REQUIRED = ["G2-01", "G2-02", "G2-06", "G2-07", "G2-08", "G2-09", "G2-10", "G2-11"]
ALL = [f"G2-{k:02d}" for k in range(1, 15)]  # the other six are optional
TESTS = {
    "G2-01": "Reference beam: every check once. Bending governs.",
    "G2-02": "Shear governs (short span, very high load): shear resistance and shear area.",
    "G2-06": "Web bearing governs: bearing resistance at the support (stiff bearing 100 mm).",
    "G2-07": "Deflection governs, variable load only, limit L/360.",
    "G2-08": "Deflection governs, total load (dead + live + self-weight, unfactored), limit L/250.",
    "G2-09": "Lateral-torsional buckling governs, restraint every 1/3 span.",
    "G2-10": "Lateral-torsional buckling governs, restraint at mid-span.",
    "G2-11": "Lateral-torsional buckling governs, no intermediate restraint.",
}
ARIAL = "Arial"
BLUE = Font(name=ARIAL, size=10, color="0000FF")
BLACK = Font(name=ARIAL, size=10)
BOLD = Font(name=ARIAL, size=10, bold=True)
WHITE = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
TITLE = Font(name=ARIAL, size=14, bold=True)
HEAD = PatternFill("solid", start_color="1F3864")
GREY = PatternFill("solid", start_color="E7E6E6")
YELLOW = PatternFill("solid", start_color="FFFF00")
LIGHT = PatternFill("solid", start_color="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
TOL = "'Read me'!$C$6"


def load(d):
    props = os.path.join(d, "UB-secpropsdimsprops-EC3UKNA-UK-10-16-2026.xlsx")
    cat, _ = parse_workbook(props)
    bend = bb.parse_bending(os.path.join(d, "UB-bending-S355-EC3UKNA-UK-10-20-2026.xlsx"))
    bear = bb.parse_bearing(os.path.join(d, "UB-webbearingbuckling-S355-EC3UKNA-UK-10-20-2026.xlsx"))
    return cat, bend, bear


def read_cases(path):
    with open(path, newline="") as f:
        rows = {r["id"]: r for r in csv.DictReader(f)}
    return [rows[i] for i in ALL]


def case_sheet(wb, case, cat, bend, bear, names_C):
    cid, name = case["id"], case["section"]
    L, w, rho = float(case["L_m"]), float(case["w_kNm"]), float(case["restraint_ratio"])
    if abs(rho - 1 / 3) < 1e-3:
        rho = 1 / 3  # the CSV stores 0.3333
    mode = "variable" if case["deflection"].startswith("variable") else "total"
    ratio = float(case["deflection"].split("L/")[1])
    i = cat.names.index(name)
    c = cat.cols
    K = Checker(cat.mass_ordered(False), replace(CheckerConfig(), restraint_ratio=rho, deflection_load=mode, deflection_ratio=ratio))
    r = K.check(L, w, K.cat.names.index(name))
    ltb = rho > 0
    ws = wb.create_sheet(cid)
    required = cid in REQUIRED
    ws.sheet_properties.tabColor = "70AD47" if required else "A6A6A6"
    desc = TESTS.get(cid) or ("Optional: " + case["label"] + ".")
    ws.sheet_view.showGridLines = False
    for col, wd in zip("ABCDEFGHI", (50, 9, 56, 15, 15, 13, 15, 13, 12)):
        ws.column_dimensions[col].width = wd
    ws["A1"] = f"{cid} ({'REQUIRED' if required else 'optional'}): {case['label']}"
    ws["A1"].font = TITLE
    ws["A2"] = f"{name}.  {desc}"
    ws["A2"].font = BOLD
    ref = {}
    row = 7

    def hdr(text):
        nonlocal row
        for col in range(1, 10):
            ws.cell(row, col).fill = HEAD
        ws.cell(row, 1, text).font = WHITE
        row += 1

    def inp(key, label, unit, src, value, fmt="General"):
        nonlocal row
        ws.cell(row, 1, label).font = BLACK
        ws.cell(row, 2, unit).font = BLACK
        ws.cell(row, 3, src).font = BLACK
        ws.cell(row, 3).alignment = WRAP
        cell = ws.cell(row, 4, value)
        cell.font, cell.number_format = BLUE, fmt
        ref[key] = f"$D${row}"
        row += 1

    hdr("INPUTS (blue = typed from the Blue Book tables or the case sheet; change nothing here)")
    inp("sec", "Section", "", "Blue Book section properties table", name)
    for key, label, unit, col, fmt in (("m", "Mass per metre", "kg/m", "m", "0.0"), ("h", "Depth h", "mm", "h", "0.0"), ("b", "Width b", "mm", "b", "0.0"),
                                       ("tw", "Web thickness tw", "mm", "tw", "0.0"), ("tf", "Flange thickness tf", "mm", "tf", "0.0"),
                                       ("r", "Root radius r", "mm", "r", "0.0"), ("A", "Area A", "cm2", "A", "0.0"),
                                       ("Iy", "Second moment of area Iy", "cm4", "Iy", "#,##0"), ("Wpl", "Plastic modulus Wpl,y", "cm3", "Wpl_y", "#,##0"),
                                       ("cf", "Flange ratio cf/tf", "", "cf_tf", "0.00"), ("cw", "Web ratio cw/tw", "", "cw_tw", "0.0")):
        inp(key, label, unit, "Blue Book section properties table", float(c[col][i]), fmt)
    inp("L", "Span L", "m", "case sheet (docs/gate2_cases.md)", L, "0.00")
    inp("w", "Factored superimposed UDL w (self-weight excluded)", "kN/m", "case sheet", w, "0.0000")
    inp("rho", "Restraint ratio Lcr/L (0 = fully restrained)", "", "case sheet", rho, "0.0000")
    inp("mode", "Deflection load", "", "case sheet: variable = live load only; total = dead + live + self-weight", mode)
    inp("ratio", "Deflection limit: span /", "", "case sheet", ratio, "0")
    inp("E", "Young's modulus E", "N/mm2", "EN 1993-1-1 3.2.6", 210000, "#,##0")
    inp("gG", "Load factor, dead and self-weight", "", "EN 1990 (assumption 2)", 1.35, "0.00")
    inp("gQ", "Load factor, live", "", "EN 1990 (assumption 2)", 1.5, "0.00")
    inp("gM0", "Partial factor gamma M0", "", "UK NA", 1.0, "0.00")
    inp("gM1", "Partial factor gamma M1", "", "UK NA", 1.0, "0.00")
    inp("phi", "Live share of the factored load (fixed, assumption 4)", "", "27.0 / 56.7675, the value of the reference case", "=27/56.7675", "0.0000")
    inp("gacc", "Weight per kg", "m/s2", "standard gravity", 9.80665, "0.00000")
    kb = bend[name]
    inp("Mc_tab", "Blue Book: Mc,y,Rd", "kNm", "Blue Book bending table", kb["Mc"], "#,##0.0")
    inp("Vc_tab", "Blue Book: Vc,Rd", "kN", "Blue Book web bearing table", bear[name]["Vc"], "#,##0.0")
    inp("F_tab", "Blue Book: FRd, ss = 100 mm, c = 0", "kN", "Blue Book web bearing table", float(bear[name]["F0"][bb.SS.index(100.0)]), "#,##0.0")
    if ltb:
        lcr = rho * L
        c1 = 1.00 if rho < 1 else 1.13
        mb = float(kb["Mb"][c1][bb.LENGTHS.index(lcr)])
        inp("Lcr", "Distance between lateral restraints Lcr", "m", "restraint ratio x span (a length listed in the Blue Book table)", lcr, "0.00")
        inp("C1", "Blue Book C1 row used", "", "1.00 where restraints lie between the supports; 1.13 for no restraint (program uses 1.127)", c1, "0.00")
        inp("Mb_tab", "Blue Book: Mb,Rd at Lcr", "kNm", "Blue Book buckling resistance moment table, row C1, column Lcr", mb, "#,##0.0")
    row += 1
    hdr("CALCULATION (black = Excel formula; grey = value printed by the program; yellow = enter your own value)")
    for col, t in enumerate(("Quantity", "Unit", "How it is calculated", "Hand calc (Excel)", "Program", "Diff hand vs program", "Your value", "Diff yours vs program", "Status"), 1):
        cell = ws.cell(row, col, t)
        cell.font, cell.fill, cell.border, cell.alignment = BOLD, LIGHT, BOX, Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 30
    row += 1
    first = row
    cr = {}

    def calc(key, label, unit, how, formula, prog, fmt="#,##0.00", fill_yours=True):
        nonlocal row
        f = formula.format(**{**ref, **cr})
        ws.cell(row, 1, label).font = BLACK
        ws.cell(row, 2, unit).font = BLACK
        ws.cell(row, 3, how).font = BLACK
        ws.cell(row, 3).alignment = WRAP
        d = ws.cell(row, 4, f)
        d.font, d.number_format = BLACK, fmt
        cr[key] = f"$D${row}"
        if prog is not None:
            e = ws.cell(row, 5, prog)
            e.font, e.number_format, e.fill = BLACK, fmt, GREY
            if not isinstance(prog, str):
                ws.cell(row, 6, f'=IF(E{row}=0,"",(D{row}-E{row})/E{row})').number_format = "0.00%"
        if fill_yours and prog is not None and not isinstance(prog, str):
            g = ws.cell(row, 7)
            g.fill, g.number_format, g.font = YELLOW, fmt, BLUE
            ws.cell(row, 8, f'=IF(G{row}="","",(G{row}-E{row})/E{row})').number_format = "0.00%"
            ws.cell(row, 9, f'=IF(G{row}="","",IF(ABS(H{row})<={TOL},"OK","CHECK"))')
        for col in range(1, 10):
            ws.cell(row, col).border = BOX
            if col in (6, 8, 9):
                ws.cell(row, col).font = BLACK
        row += 1

    calc("fy", "Yield strength fy", "N/mm2", "355, 345, 335, 325 for max(tf, tw) up to 16, 40, 63, 80 mm",
         "=IF(MAX({tf},{tw})<=16,355,IF(MAX({tf},{tw})<=40,345,IF(MAX({tf},{tw})<=63,335,325)))", r.fy, "0")
    calc("eps", "eps = sqrt(235/fy)", "", "", "=SQRT(235/{fy})", None, "0.0000")
    calc("cls", "Section class", "", "flange: 9, 10, 14 eps; web: 72, 83, 124 eps; the larger class governs",
         "=MAX(IF({cf}<={eps}*9,1,IF({cf}<={eps}*10,2,IF({cf}<={eps}*14,3,4))),IF({cw}<={eps}*72,1,IF({cw}<={eps}*83,2,IF({cw}<={eps}*124,3,4))))",
         r.section_class, "0")
    calc("hw", "Web height hw = h - 2 tf", "mm", "", "={h}-2*{tf}", None, "0.0")
    calc("g", "Self-weight g", "kN/m", "mass x 9.80665 / 1000", "={m}*{gacc}/1000", None, "0.0000")
    calc("wEd", "Total factored load wEd", "kN/m", "w + 1.35 g", "={w}+{gG}*{g}", r.w_Ed, "0.0000")
    calc("MEd", "Design moment MEd", "kNm", "wEd L^2 / 8", "={wEd}*{L}^2/8", r.M_Ed)
    calc("VEd", "Design shear force VEd", "kN", "wEd L / 2", "={wEd}*{L}/2", r.V_Ed)
    calc("Mc", "Bending resistance Mc,Rd", "kNm", "Wpl x fy / gamma M0 (classes 1 and 2)", "={Wpl}*1000*{fy}/{gM0}/10^6", r.M_c_Rd)
    calc("ub", "Bending utilisation", "", "MEd / Mc,Rd", "={MEd}/{Mc}", r.u_bending, "0.0000")
    calc("Av", "Shear area Av", "mm2", "A - 2 b tf + (tw + 2r) tf, not less than hw tw (eta = 1.0)", "=MAX({A}*100-2*{b}*{tf}+({tw}+2*{r})*{tf},{hw}*{tw})", None, "#,##0.0")
    calc("Vpl", "Shear resistance Vpl,Rd", "kN", "Av x fy / (sqrt(3) x gamma M0)", "={Av}*{fy}/SQRT(3)/{gM0}/1000", r.V_pl_Rd)
    calc("us", "Shear utilisation", "", "VEd / Vpl,Rd", "={VEd}/{Vpl}", r.u_shear, "0.0000")
    calc("sb", "Shear buckling needed?", "", "yes if hw/tw > 72 eps (not for these sections)", '=IF({hw}/{tw}>72*{eps},"yes","no")',
         "yes" if r.shear_buckling_applies else "no", "General", False)
    sb = r.shear_buckling_applies
    if sb:
        calc("lam", "Web slenderness lambda_w", "", "hw / (86.4 tw eps)  (EN 1993-1-5 5.3(5); stiffeners at the supports only)", "={hw}/(86.4*{tw}*{eps})", None, "0.0000")
        calc("chi", "Reduction factor chi_w", "", "min(1.0, 0.83 / lambda_w)  (non-rigid end post, eta = 1.0)", "=MIN(1,0.83/{lam})", None, "0.0000")
        calc("Vb", "Shear buckling resistance Vb,Rd", "kN", "chi_w x fy x hw x tw / (sqrt(3) x gamma M1); flange contribution ignored", "={chi}*{fy}*{hw}*{tw}/SQRT(3)/{gM1}/1000", r.V_b_Rd)
        calc("usb", "Shear buckling utilisation", "", "VEd / Vb,Rd", "={VEd}/{Vb}", r.u_shear_buckling, "0.0000")
    calc("Fr", "Web bearing resistance FRd", "kN", "Blue Book table value (ss = 100 mm, c = 0)", "={F_tab}", r.F_Rd)
    calc("uw", "Web bearing utilisation", "", "VEd / FRd (table value)", "={VEd}/{F_tab}", r.u_web_bearing, "0.0000")
    calc("Qk", "Live load Qk (unfactored)", "kN/m", "share x w / 1.5", "={phi}*{w}/{gQ}", None, "0.0000")
    calc("q", "Load used for deflection q", "kN/m", "variable: Qk; total: Qk + dead part (1-share) w / 1.35 + g", '=IF({mode}="variable",{Qk},{Qk}+(1-{phi})*{w}/{gG}+{g})', None, "0.0000")
    calc("dl", "Deflection", "mm", "5 q L^4 / (384 E Iy)", "=5*{q}*({L}*1000)^4/(384*{E}*{Iy}*10^4)", r.deflection)
    calc("dlim", "Deflection limit", "mm", "L / limit", "={L}*1000/{ratio}", r.deflection_limit)
    calc("ud", "Deflection utilisation", "", "deflection / limit", "={dl}/{dlim}", r.u_deflection, "0.0000")
    util_keys = ["ub", "us"] + (["usb"] if sb else []) + ["uw", "ud"]
    if ltb:
        calc("Lcr_c", "Lcr", "m", "restraint ratio x span", "={rho}*{L}", rho * L, "0.00", False)
        calc("Mb", "LTB resistance Mb,Rd", "kNm", "Blue Book table value at Lcr and C1", "={Mb_tab}", r.M_b_Rd)
        calc("ul", "LTB utilisation", "", "MEd / Mb,Rd (table value)", "={MEd}/{Mb_tab}", r.u_ltb, "0.0000")
        util_keys.append("ul")
    if sb:
        calc("Mpl", "Plastic moment resistance Mpl,Rd", "kNm", "Wpl x fy / gamma M0 (EN 1993-1-5 7.1)", "={Wpl}*1000*{fy}/{gM0}/10^6", None)
        calc("Mf", "Flange-only moment resistance Mf,Rd", "kNm", "b tf (h - tf) fy / gamma M0", "={b}*{tf}*({h}-{tf})*{fy}/{gM0}/10^6", None)
    fine = sb or r.governing == "bending_shear"
    nst, den = (51, 100) if fine else (11, 20)  # fine grid = the program's own stations (101 over the span, symmetric)
    stn = row + 6  # first station row: three rows follow (interaction, highest, governing), then a blank, a header and a title row
    calc("ui", "Bending-shear interaction (largest over stations)", "", "see the station table below; equals bending where shear does not reduce the moment resistance",
         f"=MAX($I${stn}:$I${stn + nst - 1})", r.u_bending_shear, "0.0000")
    util_keys.append("ui")
    allu = ",".join(cr[k] for k in util_keys)
    calc("umax", "Highest utilisation", "", "largest of the utilisations above", f"=MAX({allu})", r.max_utilisation, "0.0000")
    gov = r.governing
    labels = {"ub": "bending", "us": "shear", "uw": "web_bearing", "ud": "deflection", "ul": "ltb", "ui": "bending_shear", "usb": "shear_buckling"}
    nest = '"unknown"'
    for k in reversed(util_keys):
        nest = f'IF(ABS({cr[k]}-{cr["umax"]})<0.0015,"{labels[k]}",{nest})'
    calc("gov", "Governing check", "", "the check with the highest utilisation (bending and bending-shear coincide here)", "=" + nest, gov, "General", False)
    last = row - 1
    ws.cell(first - 1, 1).value = "Quantity"
    row += 1
    hdr(f"BENDING-SHEAR INTERACTION AT {nst} STATIONS ALONG HALF THE SPAN (the program uses 101 stations over the whole span; the other half is symmetric)")
    titles = ("x / L", "", "x (m)", "M(x) (kNm)", "V(x) (kN)", "V / Vb,Rd", "", "", "utilisation (EN 1993-1-5 7.1)") if sb else ("x / L", "", "x (m)", "M(x) (kNm)", "V(x) (kN)", "rho", "Mv,Rd (kNm)", "", "M / Mv,Rd")
    for col, t in enumerate(titles, 1):
        ws.cell(row, col, t).font = BOLD
    row += 1
    assert row == stn, (row, stn)
    for k in range(nst):
        rr = row + k
        ws.cell(rr, 1, f"={k}/{den}").number_format = "0.00"
        ws.cell(rr, 3, f"=A{rr}*{ref['L']}").number_format = "0.000"
        ws.cell(rr, 4, f"={cr['wEd']}*C{rr}*({ref['L']}-C{rr})/2").number_format = "#,##0.00"
        ws.cell(rr, 5, f"=ABS({cr['wEd']}*({ref['L']}/2-C{rr}))").number_format = "#,##0.00"
        if sb:
            ws.cell(rr, 6, f"=E{rr}/{cr['Vb']}").number_format = "0.0000"
            ws.cell(rr, 9, f"=IF(F{rr}>0.5,D{rr}/{cr['Mpl']}+(1-{cr['Mf']}/{cr['Mpl']})*(2*F{rr}-1)^2,D{rr}/{cr['Mc']})").number_format = "0.0000"
        else:
            ws.cell(rr, 6, f"=IF(E{rr}>0.5*{cr['Vpl']},MIN((2*E{rr}/{cr['Vpl']}-1)^2,1),0)").number_format = "0.0000"
            ws.cell(rr, 7, f"=MIN({cr['Mc']},({ref['Wpl']}*1000-F{rr}*({cr['hw']}*{ref['tw']})^2/(4*{ref['tw']}))*{cr['fy']}/{ref['gM0']}/10^6)").number_format = "#,##0.00"
            ws.cell(rr, 9, f"=D{rr}/G{rr}").number_format = "0.0000"
        for col in (1, 3, 4, 5, 6, 7, 9):
            ws.cell(rr, col).font = BLACK
    row += nst + 1
    note = ("Shear buckling applies. Utilisation = M / Mc,Rd where V <= 0.5 Vb,Rd; otherwise M / Mpl,Rd + (1 - Mf,Rd / Mpl,Rd) x (2 V / Vb,Rd - 1)^2 (EN 1993-1-5 7.1)." if sb else
            "rho = (2 V / Vpl - 1)^2 when V > 0.5 Vpl, otherwise 0 (EN 1993-1-1 6.2.8). Mv,Rd = (Wpl - rho x (hw tw)^2 / (4 tw)) fy / gamma M0, not more than Mc,Rd.")
    ws.cell(row, 1, note).font = BLACK
    ws["A3"] = "Largest difference, hand calc vs program"
    ws["D3"] = f"=MAX(MAX(F{first}:F{last}),-MIN(F{first}:F{last}))"
    ws["D3"].number_format = "0.00%"
    ws["A4"] = "Rows where you entered a value / flagged CHECK"
    ws["D4"] = f"=COUNT(G{first}:G{last})"
    ws["E4"] = f'=COUNTIF(I{first}:I{last},"CHECK")'
    ws["A5"] = "Yellow cells: enter your own value for any row. The difference and OK/CHECK appear by themselves."
    for a in ("A3", "A4", "A5"):
        ws[a].font = BOLD
    for a in ("D3", "D4", "E4"):
        ws[a].font = BOLD
    ws.freeze_panes = "A7"
    return {"id": cid, "required": required, "name": name, "tests": desc, "gov": r.governing, "umax": r.max_utilisation}


def build(d, out):
    cat, bend, bear = load(d)
    cases = read_cases(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "gate2", "cases.csv"))
    wb = Workbook()
    rd = wb.active
    rd.title = "Read me"
    summ = wb.create_sheet("Summary")
    infos = [case_sheet(wb, c, cat, bend, bear, None) for c in cases]
    rd.sheet_view.showGridLines = False
    rd.column_dimensions["A"].width = 3
    rd.column_dimensions["B"].width = 34
    rd.column_dimensions["C"].width = 100
    rd["B1"] = "Gate 2 hand check: 14 cases (8 required, 6 optional)"
    rd["B1"].font = TITLE
    rd["B3"] = "What this is"
    rd["C3"] = ("The beam program prints a utilisation (demand divided by resistance; 1.00 or less passes) for each check. This workbook lets you recompute those numbers "
                "and see the difference to the program's number. Each case has its own sheet: green tabs are the 8 required cases, grey tabs the 6 optional ones.")
    rd["B4"] = "Who does what"
    rd["C4"] = "Nothing needs to be run. Open a case sheet, enter your own value in any yellow cell, and read the difference."
    rd["B6"] = "Tolerance (editable)"
    rd["C6"] = 0.005
    rd["C6"].number_format = "0.0%"
    rd["C6"].font, rd["C6"].fill = BLUE, YELLOW
    rd["C6"].alignment = Alignment(horizontal="left")
    rd["B7"] = "Why 0.5 %"
    rd["C7"] = "The Blue Book rounds to 3 significant figures (585.75 kNm appears as 586), so differences up to about this size are rounding."
    rd["B9"] = "Colours"
    rd["C9"] = "Blue text: typed input (Blue Book or case sheet).  Black: Excel formula.  Grey fill: value printed by the program.  Yellow fill: enter your own value."
    rd["B10"] = "Columns of a case sheet"
    rd["C10"] = ("Hand calc (Excel): the quantity calculated by formula from the inputs, following the Eurocode steps.  Program: the number the checker computed.  "
                 "Diff hand vs program: (hand - program) / program.  Your value: your own result.  Diff yours vs program: (yours - program) / program.  "
                 "Status: OK if the difference is within the tolerance, CHECK otherwise.")
    rd["B11"] = "Where the program numbers come from"
    rd["C11"] = "The checker (src/ub_rl/checker.py) run for each case; they are also in data/gate2/cases.csv. The cases were selected by scripts/gate2_cases.py (see docs/gate2_cases.md)."
    rd["B12"] = "Where the Blue Book values come from"
    rd["C12"] = ("Section properties table (16/10/2026), bending table and web bearing table (20/10/2026), S355, Eurocode 3 UK NA. FRd, Mb,Rd and Vc,Rd are read from the tables; "
                 "they are not recalculated here.")
    rd["B14"] = "Example (format only)"
    rd["C14"] = "Mc,Rd: program 585.75 kNm; your value 586; difference 0.04 %; status OK."
    rd["B15"] = "What to report"
    rd["C15"] = "Any row marked CHECK, with your value and the program's value; the answers on the 12 assumptions in docs/review_request.md; any missing check."
    for rr in range(3, 16):
        rd.cell(rr, 2).font = BOLD
        rd.cell(rr, 3).alignment = WRAP
        if rr != 6:
            rd.cell(rr, 3).font = BLACK
    rd["C6"].font = BLUE
    for col, t in enumerate(("Case", "Required?", "Section", "What it tests", "Program: governing check", "Program: highest utilisation", "Largest diff, hand calc vs program", "Rows you filled", "Rows flagged CHECK", "Status"), 1):
        cell = summ.cell(1, col, t)
        cell.font, cell.fill, cell.alignment = WHITE, HEAD, Alignment(wrap_text=True, vertical="center")
    summ.row_dimensions[1].height = 42
    for col, wd in zip("ABCDEFGHIJ", (10, 11, 16, 60, 18, 16, 18, 13, 14, 22)):
        summ.column_dimensions[col].width = wd
    for k, inf in enumerate(infos, start=2):
        s = f"'{inf['id']}'"
        vals = [inf["id"], "yes" if inf["required"] else "optional", inf["name"], inf["tests"], inf["gov"], inf["umax"], f"={s}!$D$3", f"={s}!$D$4", f"={s}!$E$4",
                f'=IF(H{k}=0,"not started",IF(I{k}>0,"CHECK "&I{k}&" row(s)","all OK"))']
        for col, v in enumerate(vals, 1):
            cell = summ.cell(k, col, v)
            cell.font, cell.border = BLACK, BOX
            cell.alignment = WRAP
        summ.cell(k, 6).number_format = "0.0000"
        summ.cell(k, 7).number_format = "0.00%"
    summ.freeze_panes = "A2"
    wb.save(out)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "gate2_hand_check.xlsx"
    build("data/raw", out)
    print("written", out)
