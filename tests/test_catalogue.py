import numpy as np
import openpyxl
import pytest

from ub_rl.catalogue import CONSISTENCY, CatalogueError, parse_workbook, EXPECTED_SHA256


def test_counts_and_sha(parsed):
    cat, rep = parsed
    assert rep["sha256"] == EXPECTED_SHA256 and rep["sha256_ok"]
    assert len(cat) == 107 and rep["plus_rows"] == 16 and rep["primary_set_C"] == 91
    assert (rep["first_row"], rep["last_row"]) == (11, 117)
    assert len(set(cat.names)) == 107


def test_plus_rows(parsed):
    cat, rep = parsed
    assert sorted(rep["plus_designations"]) == sorted(
        ["610x178x100", "610x178x92", "610x178x82", "533x312x273", "533x312x219", "533x312x182", "533x312x151",
         "533x210x138", "533x165x85", "533x165x75", "533x165x66", "457x191x161", "457x191x133", "457x191x106",
         "406x178x85", "406x140x53"])
    assert not cat.mass_ordered(False).plus.any()
    assert cat.mass_ordered(True).plus.sum() == 16


def test_units_verified(parsed):
    u = parsed[1]["units"]
    assert u["m"] == "kg/m" and u["Iy"] == "cm4" and u["Wpl_y"] == "cm3" and u["A"] == "cm2" and u["h"] == "mm"


def test_mass_order(C, Cplus):
    for S in (C, Cplus):
        key = list(zip(S["m"], S["h"]))
        assert key == sorted(key)  # mass, then depth; stable sort keeps catalogue order for full ties
    assert len(C) == 91 and len(Cplus) == 107


def test_reference_section_values(C):
    i = C.index_of("457x191x74")
    assert C["m"][i] == 74.3 and C["h"][i] == 457.0 and C["tw"][i] == 9.0 and C["tf"][i] == 14.5
    assert C["Wpl_y"][i] > 1000  # cm3, order of magnitude only; exact values are used as parsed


def test_consistency_within_rounding(parsed):
    for k, dev in parsed[1]["consistency_max_abs_deviation"].items():
        assert dev <= CONSISTENCY[k], (k, dev)


def test_sha_mismatch_rejected(raw_path, tmp_path):
    p = tmp_path / "x.xlsx"
    wb = openpyxl.load_workbook(raw_path)
    wb.active["D11"] = "585.0"
    wb.save(p)
    with pytest.raises(CatalogueError, match="SHA-256"):
        parse_workbook(p)


def test_header_unit_change_rejected(raw_path, tmp_path):
    p = tmp_path / "x.xlsx"
    wb = openpyxl.load_workbook(raw_path)
    wb.active["R10"] = "mm4"
    wb.save(p)
    with pytest.raises(CatalogueError, match="header mismatch"):
        parse_workbook(p, check_sha=False)


def test_non_numeric_rejected(raw_path, tmp_path):
    p = tmp_path / "x.xlsx"
    wb = openpyxl.load_workbook(raw_path)
    wb.active["X20"] = "n/a"
    wb.save(p)
    with pytest.raises(CatalogueError, match="not numeric"):
        parse_workbook(p, check_sha=False)
