"""The Gate 2 hand-check workbook builds, has one sheet per case (14), and carries live formulas (needs data/raw workbooks)."""
import os
import pytest
import openpyxl

D = os.path.join(os.path.dirname(__file__), "..", "data", "raw")
FILES = ["UB-secpropsdimsprops-EC3UKNA-UK-10-16-2026.xlsx", "UB-bending-S355-EC3UKNA-UK-10-20-2026.xlsx", "UB-webbearingbuckling-S355-EC3UKNA-UK-10-20-2026.xlsx"]
pytestmark = pytest.mark.skipif(not all(os.path.exists(os.path.join(D, f)) for f in FILES), reason="Blue Book workbooks not in data/raw/")


def test_workbook_structure(tmp_path):
    import importlib.util
    spec = importlib.util.spec_from_file_location("mk", os.path.join(os.path.dirname(__file__), "..", "scripts", "make_hand_check_workbook.py"))
    mk = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mk)
    out = tmp_path / "w.xlsx"
    mk.build(D, str(out))
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Read me", "Summary"] + mk.ALL and len(mk.REQUIRED) == 8 and len(mk.ALL) == 14
    n_formulas = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    assert n_formulas > 1000
    ws = wb["G2-01"]
    assert ws["A1"].value.startswith("G2-01") and ws["D3"].value.startswith("=MAX(")
