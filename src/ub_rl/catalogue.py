"""Parser for the Steel for Life Blue Book UB workbook (spec section 4).

All values are read from the workbook; none are typed or corrected by hand. Column positions are
verified against the header rows (group, sub-group, symbol, unit) and the file against its SHA-256."""
import hashlib
import re
from dataclasses import dataclass, field

import numpy as np

EXPECTED_SHA256 = "190e17b694bdb27adac2a762738421590394871662236c5df2fb073231740aa8"
EXPECTED_ROWS = 107
EXPECTED_PLUS = 16
FIRST_ROW = 11

# key: (column, group label row 7, sub label row 8, symbol row 9, unit row 10)
COLUMNS = {
    "m": ("D", "Mass per metre", "", "", "kg/m"),
    "h": ("E", "Depth of section", "", "h", "mm"),
    "b": ("F", "Width of section", "", "b", "mm"),
    "tw": ("G", "Thickness", "Web", "tw", "mm"),
    "tf": ("H", "Thickness", "Flange", "tf", "mm"),
    "r": ("I", "Root radius", "", "r", "mm"),
    "d": ("J", "Depth between fillets", "", "d", "mm"),
    "cw_tw": ("K", "Ratios for local buckling", "Web", "cw / tw", ""),
    "cf_tf": ("L", "Ratios for local buckling", "Flange", "cf / tf", ""),
    "C": ("M", "Dimensions for detailing", "End clearance", "C", "mm"),
    "N": ("N", "Dimensions for detailing", "Notch", "N", "mm"),
    "n": ("O", "Dimensions for detailing", "Notch", "n", "mm"),
    "surf_m": ("P", "Surface area", "Per metre", "", "m2"),
    "surf_t": ("Q", "Surface area", "Per tonne", "", "m2"),
    "Iy": ("R", "Second moment of area", "Axis y-y", "", "cm4"),
    "Iz": ("S", "Second moment of area", "Axis z-z", "", "cm4"),
    "iy": ("T", "Radius of gyration", "Axis y-y", "", "cm"),
    "iz": ("U", "Radius of gyration", "Axis z-z", "", "cm"),
    "Wel_y": ("V", "Elastic modulus", "Axis y-y", "", "cm3"),
    "Wel_z": ("W", "Elastic modulus", "Axis z-z", "", "cm3"),
    "Wpl_y": ("X", "Plastic modulus", "Axis y-y", "", "cm3"),
    "Wpl_z": ("Y", "Plastic modulus", "Axis z-z", "", "cm3"),
    "U": ("Z", "Buckling parameter", "", "U", ""),
    "X": ("AA", "Torsional index", "", "X", ""),
    "Iw": ("AB", "Warping constant", "", "Iw", "dm6"),
    "It": ("AC", "Torsional constant", "", "IT", "cm4"),
    "A": ("AD", "Area of section", "", "A", "cm2"),
}
UNITS = {k: v[4] for k, v in COLUMNS.items()}

# Internal consistency tolerances (rounding of the published table); reported, not corrected
CONSISTENCY = {
    "d - (h - 2tf - 2r) [mm]": 0.5,
    "cw/tw - d/tw": 0.06,
    "cf/tf - (b - tw - 2r)/(2 tf)": 0.006,
    "m / (0.785 A) - 1": 0.01,
    "Wel_y / (Iy / (h/2)) - 1": 0.01,
}


class CatalogueError(Exception):
    pass


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


@dataclass
class Catalogue:
    names: list
    series: list
    plus: np.ndarray  # bool, True for rows marked "+"
    rows: np.ndarray  # worksheet row numbers
    cols: dict  # key -> float64 array (units in UNITS)
    meta: dict = field(default_factory=dict)

    def __len__(self):
        return len(self.names)

    def __getitem__(self, key):
        return self.cols[key]

    def index_of(self, name):
        return self.names.index(name)

    def subset(self, order):
        """New catalogue with the given row positions in the given order."""
        order = np.asarray(order, dtype=int)
        return Catalogue(
            names=[self.names[i] for i in order],
            series=[self.series[i] for i in order],
            plus=self.plus[order],
            rows=self.rows[order],
            cols={k: v[order] for k, v in self.cols.items()},
            meta=dict(self.meta),
        )

    def mass_ordered(self, include_plus=False):
        """Spec section 9 order: mass ascending, ties by smaller h, then catalogue order.
        include_plus=False gives the primary set C (91 rows); True gives C+ (107 rows)."""
        pos = np.arange(len(self))
        if not include_plus:
            pos = pos[~self.plus]
        key = sorted(pos, key=lambda i: (self.cols["m"][i], self.cols["h"][i], i))
        sub = self.subset(key)
        sub.meta["set"] = "C+" if include_plus else "C"
        return sub


def _merged_value(ws, cell):
    if cell.value not in (None, ""):
        return cell.value
    for rng in ws.merged_cells.ranges:
        if cell.coordinate in rng:
            return ws.cell(rng.min_row, rng.min_col).value
    return cell.value


def _norm(v):
    return "" if v is None else str(v).strip()


def _verify_header(ws):
    for key, (col, g, s, sym, unit) in COLUMNS.items():
        got = (
            _norm(_merged_value(ws, ws[f"{col}7"])),
            _norm(_merged_value(ws, ws[f"{col}8"])),
            _norm(ws[f"{col}9"].value),
            _norm(ws[f"{col}10"].value),
        )
        if got != (g, s, sym, unit):
            raise CatalogueError(f"header mismatch at column {col} ({key}): expected {(g, s, sym, unit)}, found {got}")


def _designation(series, label):
    s = re.sub(r"\s*x\s*", "x", series.strip())
    m = re.fullmatch(r"x\s*(\d+)", label.strip())
    if not m or not re.fullmatch(r"\d+x\d+", s):
        raise CatalogueError(f"cannot parse designation {series!r} / {label!r}")
    return f"{s}x{m.group(1)}"


def parse_workbook(path, check_sha=True):
    """Read the workbook into a Catalogue in worksheet order. Returns (catalogue, report)."""
    import openpyxl

    sha = sha256_file(path)
    if check_sha and sha != EXPECTED_SHA256:
        raise CatalogueError(f"SHA-256 mismatch: expected {EXPECTED_SHA256}, found {sha}")
    wb = openpyxl.load_workbook(path, data_only=True)
    if len(wb.sheetnames) != 1:
        raise CatalogueError(f"expected one sheet, found {wb.sheetnames}")
    ws = wb.active
    _verify_header(ws)

    names, series_l, plus, rows = [], [], [], []
    vals = {k: [] for k in COLUMNS}
    series = None
    r = FIRST_ROW
    while _norm(ws[f"D{r}"].value) != "":
        a, b, c = (_norm(ws[f"{x}{r}"].value) for x in "ABC")
        if a:
            series = a
        if series is None:
            raise CatalogueError(f"row {r}: no series designation")
        if c not in ("", "+"):
            raise CatalogueError(f"row {r}: unexpected marker {c!r} in column C")
        names.append(_designation(series, b))
        series_l.append(re.sub(r"\s*x\s*", "x", series))
        plus.append(c == "+")
        rows.append(r)
        for key, (col, *_rest) in COLUMNS.items():
            try:
                vals[key].append(float(ws[f"{col}{r}"].value))
            except (TypeError, ValueError):
                raise CatalogueError(f"row {r}, column {col} ({key}): not numeric: {ws[f'{col}{r}'].value!r}")
        r += 1
    last = r - 1
    footnote = _norm(ws[f"A{r + 1}"].value)
    if not footnote.startswith("+ These sections are in addition to the range of BS EN 10365"):
        raise CatalogueError(f"footnote for '+' rows not found at A{r + 1}: {footnote!r}")
    if len(set(names)) != len(names):
        raise CatalogueError("duplicate designations")
    if len(names) != EXPECTED_ROWS or sum(plus) != EXPECTED_PLUS:
        raise CatalogueError(f"expected {EXPECTED_ROWS} rows with {EXPECTED_PLUS} marked '+', found {len(names)} and {sum(plus)}")

    cat = Catalogue(
        names=names,
        series=series_l,
        plus=np.array(plus),
        rows=np.array(rows),
        cols={k: np.array(v, dtype=np.float64) for k, v in vals.items()},
        meta={"sha256": sha, "source": ws["A122"].value.strip(), "sheet": ws.title},
    )
    report = {
        "sha256": sha,
        "sha256_ok": sha == EXPECTED_SHA256,
        "rows": len(cat),
        "first_row": FIRST_ROW,
        "last_row": last,
        "plus_rows": int(cat.plus.sum()),
        "plus_designations": [n for n, p in zip(names, plus) if p],
        "primary_set_C": int((~cat.plus).sum()),
        "excluded_from_C": [n for n, p in zip(names, plus) if p],
        "units": dict(UNITS),
        "consistency_max_abs_deviation": consistency(cat),
    }
    return cat, report


def consistency(cat):
    """Largest absolute deviation of each cross-check (published values are rounded)."""
    c = cat.cols
    out = {
        "d - (h - 2tf - 2r) [mm]": c["d"] - (c["h"] - 2 * c["tf"] - 2 * c["r"]),
        "cw/tw - d/tw": c["cw_tw"] - c["d"] / c["tw"],
        "cf/tf - (b - tw - 2r)/(2 tf)": c["cf_tf"] - (c["b"] - c["tw"] - 2 * c["r"]) / 2 / c["tf"],
        "m / (0.785 A) - 1": c["m"] / (0.785 * c["A"]) - 1,
        "Wel_y / (Iy / (h/2)) - 1": c["Wel_y"] / (c["Iy"] / (c["h"] / 20)) - 1,
    }
    return {k: float(np.abs(v).max()) for k, v in out.items()}
