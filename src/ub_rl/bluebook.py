"""Gate 2: comparison of the checker with the Steel for Life Blue Book UB resistance tables (S355, EC3 UK NA).

Two downloads (20/10/2026) are used: the buckling resistance moment table (Mc,y,Rd, class, Mb,Rd by length and C1)
and the web bearing and buckling table (Vc,Rd, FRd for c = 0 by stiff bearing length ss).
The tables are rounded to 3 significant figures and are computed from unrounded section data, while the checker
reads the rounded section properties. A value therefore agrees if the table value lies within the interval spanned
by the checker result when every published input is moved by its own rounding (3 s.f. for A, Iz, It, Iw, Wpl;
0.05 mm for h, b, tw, tf, r) plus half a unit in the last table digit.

The '+' markers of these two files are not used (see docs/verification_record.md)."""
import itertools
import re

import numpy as np

from .catalogue import sha256_file
from .checker import Checker, _web_bearing
from .config import CheckerConfig

BENDING_SHA256 = "c2fedba5a2402e9668610cfce9209296744092002d7a152870c2d8bf6b4ce3dd"
BEARING_SHA256 = "fb34eee16d6dfc10f148cf7104340df4c641f2d43735683d3587a50dee6b1b92"
LENGTHS = [1.0, 1.5, 2.0, 2.5, 3.0, 3.5, 4.0, 5.0, 6.0, 7.0, 8.0, 9.0, 10.0, 11.0, 12.0, 13.0, 14.0]
C1S = [1.00, 1.13, 1.35, 1.50, 1.77, 2.00, 2.50]
SS = [0.0, 10.0, 20.0, 30.0, 40.0, 50.0, 75.0, 100.0, 150.0, 200.0, 250.0, 300.0, 350.0]


class BlueBookError(Exception):
    pass


def half_ulp(x, sf=3):
    """Half a unit in the last place of x written with sf significant figures (x may be an array or a string)."""
    x = np.abs(np.asarray(x, dtype=float))
    return 0.5 * 10.0 ** (np.floor(np.log10(np.where(x == 0, 1.0, x))) - (sf - 1))


def _designation(series, label):
    return re.sub(r"\s*x\s*", "x", series) + "x" + re.sub(r"\D", "", label)


def _blocks(ws, first_row, height):
    """Yield (row, designation, plus_marker) for each section block; series label is carried down."""
    series, r = None, first_row
    while r <= ws.max_row:
        b = ws.cell(r, 2).value
        if b and str(b).strip().startswith("x"):
            a = ws.cell(r, 1).value
            if a and str(a).strip():
                series = str(a).strip()
            yield r, _designation(series, str(b)), ws.cell(r, 3).value == "+"
            r += height
        else:
            r += 1


def parse_bending(path, check_sha=True):
    import openpyxl
    if check_sha and sha256_file(path) != BENDING_SHA256:
        raise BlueBookError("bending table: SHA-256 mismatch")
    ws = openpyxl.load_workbook(path, data_only=True).active
    if not str(ws["A3"].value).startswith("Universal beams (UB), Buckling resistance moment with S355") or ws["D5"].value != "C1(1)":
        raise BlueBookError("bending table: unexpected title or C1 column")
    if [float(ws.cell(7, c).value) for c in range(5, 22)] != LENGTHS or ws["V7"].value != "cm4":
        raise BlueBookError("bending table: unexpected lengths or Iy unit")
    out = {}
    for r, name, plus in _blocks(ws, 8, 7):
        blk = {"plus_marker": plus, "Iy": float(ws.cell(r, 22).value), "Mb": {}, "Mb_tol": {}}
        for k in range(7):
            c1 = float(ws.cell(r + k, 4).value)
            if c1 != C1S[k]:
                raise BlueBookError(f"{name}: C1 order {c1}")
            vals = [ws.cell(r + k, c).value for c in range(5, 22)]
            blk["Mb"][c1] = np.array([float(v) for v in vals])
            blk["Mb_tol"][c1] = half_ulp([float(v) for v in vals])
            t = ws.cell(r + k, 1).value
            if t and str(t).startswith("Mc,y,Rd ="):
                blk["Mc"] = float(str(t).split("=")[1])
            if t and str(t).startswith("Class ="):
                blk["cls"] = int(str(t).split("=")[1])
        if "Mc" not in blk or "cls" not in blk:
            raise BlueBookError(f"{name}: Mc,y,Rd or class missing")
        blk["Mc_tol"] = float(half_ulp(blk["Mc"]))
        out[name] = blk
    return out


def parse_bearing(path, check_sha=True):
    import openpyxl
    if check_sha and sha256_file(path) != BEARING_SHA256:
        raise BlueBookError("web bearing table: SHA-256 mismatch")
    ws = openpyxl.load_workbook(path, data_only=True).active
    if not str(ws["A3"].value).startswith("Universal beams (UB), Web bearing & buckling with S355") or ws["D8"].value != "kN":
        raise BlueBookError("web bearing table: unexpected title or unit")
    if [float(ws.cell(8, c).value) for c in range(6, 19)] != SS:
        raise BlueBookError("web bearing table: unexpected ss columns")
    out = {}
    for r, name, plus in _blocks(ws, 9, 3):
        if not (str(ws.cell(r, 5).value).startswith("FRd (c = 0)") and str(ws.cell(r + 1, 5).value).startswith("clim")):
            raise BlueBookError(f"{name}: unexpected row labels")
        vals = [ws.cell(r, c).value for c in range(6, 19)]
        out[name] = {"plus_marker": plus, "Vc": float(ws.cell(r, 4).value), "Vc_tol": float(half_ulp(ws.cell(r, 4).value)),
                     "F0": np.array([float(v) for v in vals]), "F0_tol": half_ulp([float(v) for v in vals])}
    return out


def _order(cat, tab):
    if set(cat.names) != set(tab):
        raise BlueBookError("designations of the table and the catalogue differ")
    return [tab[n] for n in cat.names]


def compare_class_mc(cat, bend):
    """Class and Mc,y,Rd, plus Iy (the published Iy is an input to the checker)."""
    K = Checker(cat, CheckerConfig())
    T = _order(cat, bend)
    cls_bad = [cat.names[i] for i, t in enumerate(T) if int(K.cls[i]) != t["cls"]]
    Wpl = cat.cols["Wpl_y"]
    lo = (Wpl - half_ulp(Wpl)) * 1e3 * K.fy / 1e6
    hi = (Wpl + half_ulp(Wpl)) * 1e3 * K.fy / 1e6
    mc_bad = [(cat.names[i], float(K.Mc[i]), t["Mc"]) for i, t in enumerate(T)
              if not (lo[i] - t["Mc_tol"] - 1e-9 <= t["Mc"] <= hi[i] + t["Mc_tol"] + 1e-9)]
    iy_bad = [cat.names[i] for i, t in enumerate(T) if abs(cat.cols["Iy"][i] - t["Iy"]) > half_ulp(t["Iy"]) * (1 + 1e-9)]
    dev = max(abs(K.Mc[i] / t["Mc"] - 1) for i, t in enumerate(T))
    return {"n": len(T), "class_mismatch": cls_bad, "Mc_outside": mc_bad, "Iy_mismatch": iy_bad, "Mc_max_rel_dev": float(dev)}


def compare_vc(cat, bear):
    K = Checker(cat, CheckerConfig())
    c = cat.cols
    N = len(cat)
    lo, hi = np.full(N, np.inf), np.full(N, -np.inf)
    for d in itertools.product((-1, 0, 1), repeat=5):
        A = (c["A"] + d[0] * half_ulp(c["A"])) * 100
        b, tf, tw, r = c["b"] + d[1] * .05, c["tf"] + d[2] * .05, c["tw"] + d[3] * .05, c["r"] + d[4] * .05
        hw = c["h"] - 2 * tf
        Av = np.maximum(A - 2 * b * tf + (tw + 2 * r) * tf, K.cfg.shear_eta * hw * tw)
        V = Av * K.fy / np.sqrt(3.0) / K.cfg.gamma_M0 / 1e3
        lo, hi = np.minimum(lo, V), np.maximum(hi, V)
    T = _order(cat, bear)
    bad = [(cat.names[i], float(K.Vpl[i]), t["Vc"]) for i, t in enumerate(T)
           if not (lo[i] - t["Vc_tol"] - 1e-9 <= t["Vc"] <= hi[i] + t["Vc_tol"] + 1e-9)]
    return {"n": len(T), "Vc_outside": bad, "Vc_max_rel_dev": float(max(abs(K.Vpl[i] / t["Vc"] - 1) for i, t in enumerate(T)))}


def compare_bearing(cat, bear, c=0.0):
    """FRd (c = 0) for every section and ss. Where ss exceeds hw the checker limits ss to hw (EN 1993-1-5 6.3(1), SCI AD 537)
    and the table does not, so only 'checker not above table' is required there."""
    cfg = CheckerConfig()
    K = Checker(cat, cfg)
    col = cat.cols
    N = len(cat)
    fy = K.fy
    hw0 = col["h"] - 2 * col["tf"]
    T = _order(cat, bear)
    strict_bad, unconservative, capped, strict_n = [], [], 0, 0
    for j, s in enumerate(SS):
        lo, hi = np.full(N, np.inf), np.full(N, -np.inf)
        for d in itertools.product((-.05, 0, .05), repeat=4):  # h, tw, tf, b
            h, tw, tf, b = col["h"] + d[0], col["tw"] + d[1], col["tf"] + d[2], col["b"] + d[3]
            F, _ = _web_bearing(h - 2 * tf, tw, tf, b, fy, s, c, cfg.gamma_M1)
            lo, hi = np.minimum(lo, F), np.maximum(hi, F)
        for i, t in enumerate(T):
            tv, tol = t["F0"][j], t["F0_tol"][j]
            if s > hw0[i]:
                capped += 1
                if lo[i] - tol - 1e-9 > tv:
                    unconservative.append((cat.names[i], s, float(lo[i]), tv))
            else:
                strict_n += 1
                if not (lo[i] - tol - 1e-9 <= tv <= hi[i] + tol + 1e-9):
                    strict_bad.append((cat.names[i], s, float(lo[i]), float(hi[i]), tv))
    return {"strict_n": strict_n, "strict_outside": strict_bad, "capped_n": capped, "capped_unconservative": unconservative,
            "min_hw": float(hw0.min())}


def compare_mb(cat, bend):
    """Mb,Rd for Lcr = tabulated length, C1 as tabulated, kc = 1 / sqrt(C1), load at the shear centre (Blue Book convention)."""
    K = Checker(cat, CheckerConfig())
    c = cat.cols
    N = len(cat)
    base = (K.Iz.copy(), K.It.copy(), K.Iw.copy(), K.W_ltb.copy())
    rel = [half_ulp(c["Iz"]) / c["Iz"], half_ulp(c["It"]) / c["It"], half_ulp(c["Iw"]) / c["Iw"], half_ulp(c["Wpl_y"]) / c["Wpl_y"]]
    T = _order(cat, bend)
    bad, total, worst_unc = [], 0, 0.0
    try:
        for C1 in C1S:
            for k, L in enumerate(LENGTHS):
                lo, hi = np.full(N, np.inf), np.full(N, -np.inf)
                for d in itertools.product((-1, 1), repeat=4):
                    K.Iz, K.It, K.Iw, K.W_ltb = (base[m] * (1 + d[m] * rel[m]) for m in range(4))
                    m = K.with_config(restraint_ratio=1.0, ltb_C1=C1, ltb_kc=1 / np.sqrt(C1))._ltb(np.arange(N), L, 1e12)[2]
                    lo, hi = np.minimum(lo, m), np.maximum(hi, m)
                for i, t in enumerate(T):
                    tv, tol = t["Mb"][C1][k], t["Mb_tol"][C1][k]
                    total += 1
                    if not (lo[i] - tol - 1e-9 <= tv <= hi[i] + tol + 1e-9):
                        bad.append((cat.names[i], C1, L, float(lo[i]), float(hi[i]), tv))
    finally:
        K.Iz, K.It, K.Iw, K.W_ltb = base
    return {"n": total, "outside": bad}
